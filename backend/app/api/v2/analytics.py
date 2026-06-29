from __future__ import annotations

from collections import defaultdict
from datetime import datetime, timedelta, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import CurrentUser, require_teacher
from app.db.models import (
    AnswerOption,
    Discipline,
    Question,
    StudentAnswer,
    TeacherDiscipline,
    TestSession,
)
from app.db.session import get_session
from app.schemas.v2 import DashboardOut, DashboardPoint, ItemStat
from app.schemas.v2 import TopicAnalyticsBrief, TopicAnalyticsOut
from app.services.topic_analytics import build_topic_analytics, list_topic_analytics_briefs

router = APIRouter(prefix="/teacher/analytics", tags=["v2.analytics"])


@router.get("/topics", response_model=list[TopicAnalyticsBrief])
async def topic_analytics_topics(
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
    discipline_id: Optional[int] = None,
) -> list[TopicAnalyticsBrief]:
    return await list_topic_analytics_briefs(
        session,
        teacher_id=user.id,
        discipline_id=discipline_id,
    )


@router.get("/topics/{topic_id}", response_model=TopicAnalyticsOut)
async def topic_analytics_detail(
    topic_id: int,
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
) -> TopicAnalyticsOut:
    try:
        return await build_topic_analytics(session, teacher_id=user.id, topic_id=topic_id)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail="topic not found") from exc


@router.get("/dashboard", response_model=DashboardOut)
async def dashboard(
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
    range: str = Query("30d", pattern="^(7d|30d|90d|all)$"),
    discipline_id: Optional[int] = None,
):
    days_map = {"7d": 7, "30d": 30, "90d": 90, "all": 3650}
    since = datetime.now(timezone.utc) - timedelta(days=days_map[range])
    my_discs = (await session.execute(
        select(Discipline.discipline_id).join(
            TeacherDiscipline, TeacherDiscipline.discipline_id == Discipline.discipline_id
        ).where(TeacherDiscipline.teacher_id == user.id)
    )).scalars().all()
    if not my_discs:
        return DashboardOut(series=[], items=[], avg_overall=0.0, attempts_total=0)
    if discipline_id is not None and discipline_id in my_discs:
        my_discs = [discipline_id]
    base_q = (
        select(TestSession)
        .where(TestSession.teacher_id == user.id,
               TestSession.discipline_id.in_(my_discs),
               TestSession.completed_at.isnot(None))
    )
    if range != "all":
        base_q = base_q.where(TestSession.started_at >= since)
    sessions = (await session.execute(base_q)).scalars().all()
    if not sessions:
        return DashboardOut(series=[], items=[], avg_overall=0.0, attempts_total=0)

    daily: dict[str, list[float]] = defaultdict(list)
    total_correct = 0.0
    total_max = 0
    for s in sessions:
        pct = (s.score * 100.0 / s.max_score) if s.max_score else 0.0
        d = s.started_at.date().isoformat()
        daily[d].append(pct)
        total_correct += s.score
        total_max += s.max_score

    series = [
        DashboardPoint(date=k, avg_score=round(sum(v) / len(v), 2), attempts_count=len(v))
        for k, v in sorted(daily.items())
    ]
    avg_overall = round(total_correct * 100.0 / total_max, 2) if total_max else 0.0

    from sqlalchemy import case as sqlcase
    qs = (await session.execute(
        select(
            StudentAnswer.question_id,
            func.count().label("cnt"),
            func.sum(sqlcase((AnswerOption.is_correct == True, 1), else_=0)).label("correct_cnt"),
            Question.text,
        )
        .join(TestSession, TestSession.session_id == StudentAnswer.session_id)
        .join(AnswerOption, AnswerOption.option_id == StudentAnswer.selected_option_id)
        .join(Question, Question.question_id == StudentAnswer.question_id)
        .where(TestSession.teacher_id == user.id,
               TestSession.discipline_id.in_(my_discs),
               TestSession.completed_at.isnot(None),
               StudentAnswer.selected_option_id.isnot(None))
        .group_by(StudentAnswer.question_id, Question.text)
    )).all()
    items = [
        ItemStat(
            question_id=qid, question_text=qt,
            correct_pct=round((cc or 0) * 100.0 / cnt, 1),
        ) for qid, cnt, cc, qt in qs if cnt
    ]
    items.sort(key=lambda x: x.correct_pct)

    return DashboardOut(
        series=series, items=items,
        avg_overall=avg_overall,
        attempts_total=len(sessions),
    )

@router.get("/topics/{topic_id}/export.xlsx")
async def export_topic_analytics(
    topic_id: int,
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
):
    import io
    from datetime import datetime
    from urllib.parse import quote
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    try:
        data = await build_topic_analytics(session, teacher_id=user.id, topic_id=topic_id)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail="topic not found") from exc

    wb = Workbook()
    
    # ---- Sheet 1: Сводные показатели ----
    ws1 = wb.active
    ws1.title = "Сводные показатели"
    
    title_font = Font(name="Segoe UI", size=14, bold=True, color="1F2937")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Segoe UI", size=11, bold=True)
    regular_font = Font(name="Segoe UI", size=11)
    
    header_fill = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
    
    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB")
    )
    
    # Title
    ws1.cell(row=1, column=1, value=f"Аналитика темы: {data.summary.topic_name}").font = title_font
    ws1.cell(row=2, column=1, value=f"Дисциплина: {data.summary.discipline_name}").font = regular_font
    ws1.cell(row=3, column=1, value=f"Выгружено: {datetime.now().strftime("%d.%m.%Y %H:%M")}").font = regular_font
    
    # KPI Table
    ws1.cell(row=5, column=1, value="Показатель").font = header_font
    ws1.cell(row=5, column=1).fill = header_fill
    ws1.cell(row=5, column=2, value="Значение").font = header_font
    ws1.cell(row=5, column=2).fill = header_fill
    
    kpis = [
        ("Средний балл %", f"{data.summary.avg_score_percent or 0}%"),
        ("Медиана %", f"{data.summary.median_score_percent or 0}%"),
        ("Процент сдавших %", f"{data.summary.pass_rate_percent or 0}%"),
        ("Минимум %", f"{data.summary.min_score_percent or 0}%"),
        ("Максимум %", f"{data.summary.max_score_percent or 0}%"),
        ("Всего попыток", data.summary.attempts_total),
        ("Всего студентов", data.summary.students_total),
    ]
    for idx, (label, val) in enumerate(kpis, start=6):
        c1 = ws1.cell(row=idx, column=1, value=label)
        c2 = ws1.cell(row=idx, column=2, value=val)
        c1.font = regular_font
        c2.font = bold_font
        c1.border = thin_border
        c2.border = thin_border
        c2.alignment = Alignment(horizontal="right")
        
    # Groups Table
    start_row_groups = 15
    ws1.cell(row=start_row_groups, column=1, value="Успеваемость по группам").font = Font(name="Segoe UI", size=12, bold=True)
    
    group_headers = ["Группа", "Кол-во попыток", "Кол-во студентов", "Средний балл %", "% Сдавших"]
    for col_idx, h in enumerate(group_headers, start=1):
        cell = ws1.cell(row=start_row_groups+1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    for idx, g in enumerate(data.groups, start=start_row_groups+2):
        row_vals = [
            g.group_name,
            g.attempts_count,
            g.students_count,
            g.avg_score_percent if g.avg_score_percent is not None else "—",
            g.pass_rate_percent if g.pass_rate_percent is not None else "—"
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws1.cell(row=idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if col_idx > 1:
                cell.alignment = Alignment(horizontal="right")
                
    # Add Bar Chart for Groups
    if data.groups:
        chart_groups = BarChart()
        chart_groups.type = "col"
        chart_groups.style = 10
        chart_groups.title = "Средний балл по группам (%)"
        chart_groups.y_axis.title = "%"
        chart_groups.x_axis.title = "Группа"
        
        data_ref = Reference(ws1, min_col=4, min_row=start_row_groups+1, max_row=start_row_groups+1+len(data.groups))
        cats_ref = Reference(ws1, min_col=1, min_row=start_row_groups+2, max_row=start_row_groups+1+len(data.groups))
        chart_groups.add_data(data_ref, titles_from_data=True)
        chart_groups.set_categories(cats_ref)
        chart_groups.legend = None
        chart_groups.width = 15
        chart_groups.height = 10
        ws1.add_chart(chart_groups, "G5")

    # ---- Sheet 2: Динамика и Распределение ----
    ws2 = wb.create_sheet(title="Динамика и распределение")
    
    # Trend Table
    ws2.cell(row=1, column=1, value="Динамика результатов").font = Font(name="Segoe UI", size=12, bold=True)
    trend_headers = ["Дата", "Попыток", "Средний балл %"]
    for col_idx, h in enumerate(trend_headers, start=1):
        cell = ws2.cell(row=2, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        
    for idx, t in enumerate(data.trend, start=3):
        row_vals = [t.date, t.attempts_count, t.avg_score_percent if t.avg_score_percent is not None else 0]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws2.cell(row=idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if col_idx > 1:
                cell.alignment = Alignment(horizontal="right")
                
    # Add Line Chart for Trend
    if data.trend:
        chart_trend = LineChart()
        chart_trend.title = "Динамика успеваемости"
        chart_trend.style = 13
        chart_trend.y_axis.title = "Средний %"
        chart_trend.x_axis.title = "Дата"
        
        data_ref = Reference(ws2, min_col=3, min_row=2, max_row=2+len(data.trend))
        cats_ref = Reference(ws2, min_col=1, min_row=3, max_row=2+len(data.trend))
        chart_trend.add_data(data_ref, titles_from_data=True)
        chart_trend.set_categories(cats_ref)
        chart_trend.legend = None
        chart_trend.width = 16
        chart_trend.height = 10
        ws2.add_chart(chart_trend, "E2")
        
    # Distribution Table
    start_row_dist = max(len(data.trend) + 5, 20)
    ws2.cell(row=start_row_dist, column=1, value="Распределение оценок").font = Font(name="Segoe UI", size=12, bold=True)
    dist_headers = ["Диапазон %", "Количество"]
    for col_idx, h in enumerate(dist_headers, start=1):
        cell = ws2.cell(row=start_row_dist+1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        
    for idx, d in enumerate(data.distribution, start=start_row_dist+2):
        row_vals = [d.label, d.count]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws2.cell(row=idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if col_idx > 1:
                cell.alignment = Alignment(horizontal="right")
                
    # Add Bar Chart for Distribution
    if data.distribution:
        chart_dist = BarChart()
        chart_dist.type = "col"
        chart_dist.style = 11
        chart_dist.title = "Распределение баллов"
        chart_dist.y_axis.title = "Количество"
        chart_dist.x_axis.title = "Диапазон %"
        
        data_ref = Reference(ws2, min_col=2, min_row=start_row_dist+1, max_row=start_row_dist+1+len(data.distribution))
        cats_ref = Reference(ws2, min_col=1, min_row=start_row_dist+2, max_row=start_row_dist+1+len(data.distribution))
        chart_dist.add_data(data_ref, titles_from_data=True)
        chart_dist.set_categories(cats_ref)
        chart_dist.legend = None
        chart_dist.width = 15
        chart_dist.height = 10
        ws2.add_chart(chart_dist, f"E{start_row_dist}")

    # ---- Sheet 3: Сложные вопросы ----
    ws3 = wb.create_sheet(title="Сложные вопросы")
    ws3.cell(row=1, column=1, value="Статистика ответов на вопросы темы (Топ-20 сложных)").font = Font(name="Segoe UI", size=12, bold=True)
    
    q_headers = ["ID Вопроса", "Текст вопроса", "Тип", "Попыток", "% Верных", "Ошибок"]
    for col_idx, h in enumerate(q_headers, start=1):
        cell = ws3.cell(row=2, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        
    for idx, q in enumerate(data.difficult_questions, start=3):
        row_vals = [
            q.question_id,
            q.question_text,
            q.qtype,
            q.attempts_count,
            q.correct_percent if q.correct_percent is not None else "—",
            q.wrong_count
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws3.cell(row=idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if col_idx in [1, 3]:
                cell.alignment = Alignment(horizontal="center")
            elif col_idx > 3:
                cell.alignment = Alignment(horizontal="right")
                
    # Auto-adjust columns for all sheets
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = 0
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 70)
            
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    
    encoded_fname = quote(f"analytics-topic-{topic_id}.xlsx")
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename=\"analytics-topic-{topic_id}.xlsx\"; filename*=UTF-8''{encoded_fname}"
        }
    )
