from __future__ import annotations

import csv
import io
from typing import List, Optional

from pydantic import BaseModel, Field
from fastapi import APIRouter, Depends, HTTPException, Query, Request, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import CurrentUser, require_teacher
from app.db.models import (
    AnswerOption,
    DisciplineTopic,
    Question,
    QuestionAcceptableAnswer,
    QuestionClozeBlank,
    QuestionImage,
    QuestionTag,
    QuestionTagMap,
    StudentAnswer,
    TestSession,
)
from app.db.session import get_session
from app.schemas.v2 import (
    OptionV2,
    QuestionImportError,
    QuestionImportResult,
    QuestionOutV2,
    QuestionQualityReportOut,
)
from app.services.audit_trail import client_ip
from app.services.question_images import delete_storage_file, image_out
from app.services.question_quality import build_question_quality_report
from app.services.teacher_audit import audit as teacher_audit_log
from app.services.versioning import ensure_question_version

router = APIRouter(prefix="/teacher/questions", tags=["v2.questions"])


def _own_disciplines_subq(teacher_id: int):
    return select(Question.discipline_id).where(
        Question.discipline_id.in_(
            select(text("discipline_id")).select_from(
                text("teacher_disciplines")
            ).where(text("teacher_id = :tid")).bindparams(tid=teacher_id)
        )
    )


async def _ensure_own(session: AsyncSession, teacher_id: int, discipline_id: int) -> None:
    n = (await session.execute(
        text("SELECT COUNT(*) FROM teacher_disciplines WHERE teacher_id=:t AND discipline_id=:d")
        .bindparams(t=teacher_id, d=discipline_id)
    )).scalar_one()
    if not n:
        raise HTTPException(status_code=403, detail="not your discipline")


async def _ensure_topic_for_question(
    session: AsyncSession,
    discipline_id: int,
    topic_id: int | None,
) -> DisciplineTopic | None:
    if topic_id is None:
        return None
    topic = (await session.execute(
        select(DisciplineTopic).where(DisciplineTopic.topic_id == topic_id)
    )).scalar_one_or_none()
    if topic is None or topic.archived_at is not None:
        raise HTTPException(status_code=404, detail="topic not found")
    if topic.discipline_id != discipline_id:
        raise HTTPException(status_code=400, detail="topic does not belong to discipline")
    return topic


def _question_snapshot(q: Question) -> dict:
    return {
        "question_id": q.question_id,
        "discipline_id": q.discipline_id,
        "topic_id": q.topic_id,
        "text": q.text,
        "difficulty": q.difficulty,
        "qtype": q.qtype,
        "archived_at": q.archived_at,
    }


@router.get("", response_model=List[QuestionOutV2])
async def list_questions(
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
    q: Optional[str] = None,
    qtype: Optional[str] = None,
    tag: Optional[str] = None,
    discipline_id: Optional[int] = None,
    topic_id: Optional[int] = None,
    topic: Optional[str] = None,
    include_archived_topics: bool = False,
    archived: bool = False,
    difficulty: Optional[int] = None,
    ai_status: Optional[str] = None,
):
    stmt = select(Question).where(
        Question.discipline_id.in_(
            text("SELECT discipline_id FROM teacher_disciplines WHERE teacher_id = :tid").bindparams(tid=user.id)
        )
    )
    if archived:
        stmt = stmt.where(Question.archived_at.is_not(None))
    else:
        stmt = stmt.where(Question.archived_at.is_(None))
    if q:
        stmt = stmt.where(Question.text.ilike(f"%{q}%"))
    if qtype:
        stmt = stmt.where(Question.qtype == qtype)
    if difficulty is not None:
        stmt = stmt.where(Question.difficulty == difficulty)
    if ai_status is not None and ai_status != "all":
        if ai_status in ("manual", "none"):
            stmt = stmt.where(Question.ai_status.is_(None))
        else:
            stmt = stmt.where(Question.ai_status == ai_status)
    if discipline_id is not None:
        stmt = stmt.where(Question.discipline_id == discipline_id)
    if topic_id is not None:
        topic_row = (await session.execute(
            select(DisciplineTopic).where(DisciplineTopic.topic_id == topic_id)
        )).scalar_one_or_none()
        if topic_row is None:
            raise HTTPException(status_code=404, detail="topic not found")
        await _ensure_own(session, user.id, topic_row.discipline_id)
        stmt = stmt.where(Question.topic_id == topic_id)
    elif topic == "none":
        stmt = stmt.where(Question.topic_id.is_(None))
    if not include_archived_topics:
        stmt = stmt.where(
            (Question.topic_id.is_(None)) |
            (Question.topic_id.in_(
                select(DisciplineTopic.topic_id).where(DisciplineTopic.archived_at.is_(None))
            ))
        )
    stmt = stmt.order_by(Question.discipline_id, Question.question_id)


    questions = (await session.execute(stmt)).scalars().all()
    if not questions:
        return []
    q_ids = [q.question_id for q in questions]
    opts = (await session.execute(
        select(AnswerOption).where(AnswerOption.question_id.in_(q_ids)).order_by(
            AnswerOption.question_id, AnswerOption.option_number
        )
    )).scalars().all()
    grouped: dict[int, list[AnswerOption]] = {}
    for o in opts:
        grouped.setdefault(o.question_id, []).append(o)
    tag_rows = (await session.execute(
        select(QuestionTagMap.question_id, QuestionTag.name)
        .join(QuestionTag, QuestionTag.tag_id == QuestionTagMap.tag_id)
        .where(QuestionTagMap.question_id.in_(q_ids))
    )).all()
    tags_by_q: dict[int, list[str]] = {}
    for qid, tname in tag_rows:
        tags_by_q.setdefault(qid, []).append(tname)

    if tag:
        questions = [q for q in questions if tag in tags_by_q.get(q.question_id, [])]
        q_ids = [q.question_id for q in questions]

    images_by_q: dict[int, QuestionImage] = {}
    if q_ids:
        image_rows = (await session.execute(
            select(QuestionImage).where(QuestionImage.question_id.in_(q_ids))
        )).scalars().all()
        images_by_q = {img.question_id: img for img in image_rows}
    topic_names: dict[int, str] = {}
    topic_ids = sorted({qq.topic_id for qq in questions if qq.topic_id is not None})
    if topic_ids:
        topic_rows = (await session.execute(
            select(DisciplineTopic.topic_id, DisciplineTopic.name).where(DisciplineTopic.topic_id.in_(topic_ids))
        )).all()
        topic_names = {tid: name for tid, name in topic_rows}

    acceptable_by_q: dict[int, list[str]] = {qid: [] for qid in q_ids}
    if q_ids:
        aa_rows = (await session.execute(
            text("SELECT question_id, answer FROM question_acceptable_answers WHERE question_id = ANY(:ids) ORDER BY question_id, ord")
            .bindparams(ids=q_ids)
        )).all()
        for qid, ans in aa_rows:
            acceptable_by_q.setdefault(qid, []).append(ans)

    cloze_by_q: dict[int, list[dict]] = {qid: [] for qid in q_ids}
    if q_ids:
        cloze_rows = (await session.execute(
            select(QuestionClozeBlank).where(QuestionClozeBlank.question_id.in_(q_ids)).order_by(
                QuestionClozeBlank.question_id, QuestionClozeBlank.blank_index
            )
        )).scalars().all()
        for b in cloze_rows:
            cloze_by_q.setdefault(b.question_id, []).append({
                "index": b.blank_index,
                "kind": b.kind,
                "options": b.options,
                "correct_index": b.correct_index,
                "acceptable_answers": b.acceptable_answers,
                "case_sensitive": b.case_sensitive,
                "trim_whitespace": b.trim_whitespace,
                "normalize_universal": b.normalize_universal,
            })

    return [
        QuestionOutV2(
            question_id=qq.question_id,
            discipline_id=qq.discipline_id,
            topic_id=qq.topic_id,
            topic_name=topic_names.get(qq.topic_id) if qq.topic_id is not None else None,
            text=qq.text,
            difficulty=qq.difficulty,
            qtype=qq.qtype,
            options=[
                OptionV2(
                    option_number=o.option_number,
                    text=o.text,
                    is_correct=o.is_correct,
                    match_left=o.match_left,
                    match_right=o.match_right,
                )
                for o in grouped.get(qq.question_id, [])
            ],
            tags=tags_by_q.get(qq.question_id, []),
            short_pattern=qq.short_pattern,
            numeric_tolerance=float(qq.numeric_tolerance) if qq.numeric_tolerance is not None else None,
            match_pairs=qq.match_pairs or [],
            image=image_out(images_by_q[qq.question_id]) if qq.question_id in images_by_q else None,
            archived=qq.archived_at is not None,
            correct_bool=qq.correct_bool,
            explanation=qq.explanation,
            case_sensitive=qq.case_sensitive,
            trim_whitespace=qq.trim_whitespace,
            normalize_universal=qq.normalize_universal,
            text_mode=qq.text_mode,
            allow_partial=qq.allow_partial,
            acceptable_answers=acceptable_by_q.get(qq.question_id, []),
            cloze_blanks=cloze_by_q.get(qq.question_id, []),
            created_at=qq.created_at,
            archived_at=qq.archived_at,
            ai_status=qq.ai_status,
            ai_model_used=qq.ai_model_used,
            ai_reviewed_by=qq.ai_reviewed_by,
            ai_reviewed_at=qq.ai_reviewed_at,
            ai_generated_at=qq.ai_generated_at,
        )
        for qq in questions
    ]


@router.post("/import", response_model=QuestionImportResult)
async def import_questions(
    request: Request,
    file: UploadFile = File(...),
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
):
    raw = await file.read()
    filename = (file.filename or "").lower()
    
    rows = []
    if filename.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        headers = [str(h).strip().lower() for h in next(iterator, []) if h is not None]
        for idx, values in enumerate(iterator, start=1):
            row = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))}
            if any(str(v).strip() for v in row.values() if v is not None):
                rows.append((idx, row))
    else:
        text_stream = io.StringIO(raw.decode("utf-8-sig", errors="replace"))
        reader = csv.DictReader(text_stream)
        for idx, row in enumerate(reader, start=1):
            clean_row = {str(k).strip().lower(): v for k, v in row.items() if k is not None}
            if any(str(v).strip() for v in clean_row.values() if v is not None):
                rows.append((idx, clean_row))

    result = QuestionImportResult(created=0, skipped=0, errors=[])
    created_ids: list[int] = []
    for idx, row in rows:
        try:
            discipline_id = int(float(row.get("discipline_id") or 0))
            if not discipline_id:
                raise ValueError("discipline_id is required")
            await _ensure_own(session, user.id, discipline_id)
            topic_id: int | None = None
            if row.get("topic_id"):
                topic_id = int(float(row["topic_id"]))
            elif row.get("topic_name"):
                topic_match = (await session.execute(
                    select(DisciplineTopic).where(
                        DisciplineTopic.discipline_id == discipline_id,
                        DisciplineTopic.name == str(row["topic_name"]).strip(),
                        DisciplineTopic.archived_at.is_(None),
                    )
                )).scalar_one_or_none()
                if topic_match is None:
                    raise ValueError("topic not found")
                topic_id = topic_match.topic_id
            await _ensure_topic_for_question(session, discipline_id, topic_id)
            qtype = str(row.get("qtype") or "single").strip()
            if qtype not in {"single", "multi", "short", "numeric", "match"}:
                raise ValueError(f"invalid qtype: {qtype}")
            q = Question(
                discipline_id=discipline_id,
                topic_id=topic_id,
                text=str(row.get("text") or "").strip(),
                difficulty=int(float(row.get("difficulty") or 1)),
                qtype=qtype,
                short_pattern=str(row.get("short_pattern") or "").strip() or None,
                numeric_tolerance=float(row["numeric_tolerance"]) if row.get("numeric_tolerance") else None,
            )
            session.add(q)
            await session.flush()
            for i in range(1, 5):
                otext = str(row.get(f"option_{i}") or "").strip()
                if not otext:
                    continue
                omark = str(row.get(f"option_{i}_correct") or "").strip().lower() in {"1", "true", "yes", "y"}
                session.add(AnswerOption(
                    question_id=q.question_id,
                    option_number=i,
                    text=otext,
                    is_correct=omark,
                    match_left=str(row.get(f"option_{i}_left") or "").strip() or None,
                    match_right=str(row.get(f"option_{i}_right") or "").strip() or None,
                ))
            tag_names = [t.strip() for t in str(row.get("tags") or "").split(",") if t.strip()]
            for tname in tag_names:
                t = (await session.execute(
                    select(QuestionTag).where(QuestionTag.name == tname)
                )).scalar_one_or_none()
                if not t:
                    t = QuestionTag(name=tname)
                    session.add(t)
                    await session.flush()
                session.add(QuestionTagMap(question_id=q.question_id, tag_id=t.tag_id))
            await ensure_question_version(session, q.question_id, created_by_teacher_id=user.id)
            result.created += 1
            created_ids.append(q.question_id)
        except Exception as e:
            result.errors.append(QuestionImportError(row=idx, message=str(e)))
    if created_ids:
        await teacher_audit_log(
            session,
            teacher_id=user.id,
            action="questions_imported",
            target_type="question",
            target_id=None,
            before=None,
            after={
                "created": result.created,
                "question_ids": created_ids,
                "errors": len(result.errors),
                "filename": file.filename,
            },
            ip_addr=client_ip(request),
        )
    await session.commit()
    return result


@router.get("/export")
async def export_questions(
    format: str = Query(default="csv", pattern="^(csv|xlsx)$"),
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
    qtype: Optional[str] = None,
    tag: Optional[str] = None,
    discipline_id: Optional[int] = None,
):
    qs = await list_questions(user=user, session=session, qtype=qtype, tag=tag, discipline_id=discipline_id)
    headers_list = [
        "discipline_id", "topic_id", "topic_name", "qtype", "text", "difficulty", "short_pattern",
        "numeric_tolerance", "tags",
        "image_original_name", "image_url",
        "option_1", "option_1_correct", "option_1_left", "option_1_right",
        "option_2", "option_2_correct", "option_2_left", "option_2_right",
        "option_3", "option_3_correct", "option_3_left", "option_3_right",
        "option_4", "option_4_correct", "option_4_left", "option_4_right",
    ]

    if format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Вопросы"

        # Write header
        ws.append(headers_list)

        # Style header (Dark Gray, Bold White Text)
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_num in range(1, len(headers_list) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # Add data rows
        for q in qs:
            opts = {o.option_number: o for o in q.options}
            row = [
                q.discipline_id, q.topic_id or "", q.topic_name or "", q.qtype, q.text, q.difficulty, q.short_pattern or "",
                q.numeric_tolerance if q.numeric_tolerance is not None else "",
                ",".join(q.tags),
                q.image.original_name if q.image else "",
                q.image.url if q.image else ""
            ]
            for i in range(1, 5):
                o = opts.get(i)
                if o:
                    row += [o.text, 1 if o.is_correct else 0,
                            o.match_left or "", o.match_right or ""]
                else:
                    row += ["", "", "", ""]
            ws.append(row)

        # Borders and alignments
        thin_border = Border(
            left=Side(style="thin", color="E5E7EB"),
            right=Side(style="thin", color="E5E7EB"),
            top=Side(style="thin", color="E5E7EB"),
            bottom=Side(style="thin", color="E5E7EB")
        )
        for r in range(2, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = thin_border
                # Center short values
                if headers_list[c-1] in ["discipline_id", "topic_id", "qtype", "difficulty", "numeric_tolerance", "option_1_correct", "option_2_correct", "option_3_correct", "option_4_correct"]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-adjust column widths
        ws.row_dimensions[1].height = 28
        for col in ws.columns:
            max_len = 0
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            iter([buf.getvalue()]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="questions.xlsx"'},
        )

    # Fallback to CSV
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers_list)
    for q in qs:
        opts = {o.option_number: o for o in q.options}
        row = [q.discipline_id, q.topic_id or "", q.topic_name or "", q.qtype, q.text, q.difficulty, q.short_pattern or "",
               q.numeric_tolerance if q.numeric_tolerance is not None else "",
               ",".join(q.tags),
               q.image.original_name if q.image else "",
               q.image.url if q.image else ""]
        for i in range(1, 5):
            o = opts.get(i)
            if o:
                row += [o.text, "1" if o.is_correct else "0",
                        o.match_left or "", o.match_right or ""]
            else:
                row += ["", "", "", ""]
        w.writerow(row)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="questions.csv"'},
    )


@router.get("/quality", response_model=QuestionQualityReportOut)
async def question_quality(
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
    discipline_id: Optional[int] = Query(default=None),
    topic_id: Optional[int] = Query(default=None),
    topic: Optional[str] = Query(default=None),
    archived: bool = False,
    include_archived: bool = False,
):
    if topic is not None and topic != "none":
        raise HTTPException(status_code=400, detail="unsupported topic filter")
    if discipline_id is not None:
        await _ensure_own(session, user.id, discipline_id)
    if topic_id is not None:
        topic_row = (await session.execute(
            select(DisciplineTopic).where(DisciplineTopic.topic_id == topic_id)
        )).scalar_one_or_none()
        if topic_row is None:
            raise HTTPException(status_code=404, detail="topic not found")
        await _ensure_own(session, user.id, topic_row.discipline_id)
        if discipline_id is not None and topic_row.discipline_id != discipline_id:
            raise HTTPException(status_code=400, detail="topic does not belong to discipline")

    report = await build_question_quality_report(
        session,
        teacher_id=user.id,
        discipline_id=discipline_id,
        topic_id=topic_id,
        topic_none=topic == "none",
        archived_only=archived,
        include_archived=include_archived,
    )
    return QuestionQualityReportOut.model_validate(report, from_attributes=True)


def _now_utc():
    from datetime import datetime, timezone
    return datetime.now(timezone.utc)


@router.delete("/{question_id}", status_code=204)
async def delete_question(
    question_id: int,
    request: Request,
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
):
    q = (await session.execute(
        select(Question).where(
            Question.question_id == question_id,
            Question.archived_at.is_(None),
        )
    )).scalar_one_or_none()
    if q is None:
        raise HTTPException(status_code=404, detail="not found")
    await _ensure_own(session, user.id, q.discipline_id)
    before = _question_snapshot(q)
    used = (await session.execute(
        select(StudentAnswer).where(StudentAnswer.question_id == question_id).limit(1)
    )).first()
    if used:
        sessions = (await session.execute(
            select(TestSession.session_id).where(TestSession.completed_at.is_(None)).limit(1)
        )).first()
        if sessions:
            raise HTTPException(status_code=409, detail="question used in attempts")
    q.archived_at = _now_utc()
    await ensure_question_version(session, q.question_id, created_by_teacher_id=user.id)
    await teacher_audit_log(
        session,
        teacher_id=user.id,
        action="question_archived",
        target_type="question",
        target_id=question_id,
        before=before,
        after=_question_snapshot(q),
        ip_addr=client_ip(request),
    )
    await session.commit()
    return None


@router.post("/{question_id}/restore", response_model=QuestionOutV2)
async def restore_question(
    question_id: int,
    request: Request,
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
):
    q = (await session.execute(
        select(Question).where(Question.question_id == question_id)
    )).scalar_one_or_none()
    if q is None:
        raise HTTPException(status_code=404, detail="not found")
    await _ensure_own(session, user.id, q.discipline_id)
    before = _question_snapshot(q)
    if q.archived_at is None:
        # Even if already active, return the row for symmetry.
        topic_row = (await session.execute(
            select(DisciplineTopic).where(DisciplineTopic.topic_id == q.topic_id)
        )).scalar_one_or_none() if q.topic_id else None
        return QuestionOutV2(
            question_id=q.question_id,
            discipline_id=q.discipline_id,
            topic_id=q.topic_id,
            topic_name=topic_row.name if topic_row else None,
            text=q.text,
            difficulty=q.difficulty,
            qtype=q.qtype,
            options=[
                OptionV2(
                    option_number=o.option_number,
                    text=o.text,
                    is_correct=o.is_correct,
                    match_left=o.match_left,
                    match_right=o.match_right,
                ) for o in q.options
            ],
            tags=[],
            short_pattern=q.short_pattern,
            numeric_tolerance=float(q.numeric_tolerance) if q.numeric_tolerance is not None else None,
            match_pairs=q.match_pairs or [],
            image=None,
            archived=False,
        )
    q.archived_at = None
    await ensure_question_version(session, q.question_id, created_by_teacher_id=user.id)
    await teacher_audit_log(
        session,
        teacher_id=user.id,
        action="question_restored",
        target_type="question",
        target_id=question_id,
        before=before,
        after=_question_snapshot(q),
        ip_addr=client_ip(request),
    )
    await session.commit()
    await session.refresh(q)
    topic_row = (await session.execute(
        select(DisciplineTopic).where(DisciplineTopic.topic_id == q.topic_id)
    )).scalar_one_or_none() if q.topic_id else None
    return QuestionOutV2(
        question_id=q.question_id,
        discipline_id=q.discipline_id,
        topic_id=q.topic_id,
        topic_name=topic_row.name if topic_row else None,
        text=q.text,
        difficulty=q.difficulty,
        qtype=q.qtype,
        options=[
            OptionV2(
                option_number=o.option_number,
                text=o.text,
                is_correct=o.is_correct,
                match_left=o.match_left,
                match_right=o.match_right,
            ) for o in q.options
        ],
        tags=[],
        short_pattern=q.short_pattern,
        numeric_tolerance=float(q.numeric_tolerance) if q.numeric_tolerance is not None else None,
        match_pairs=q.match_pairs or [],
        image=None,
        archived=False,
    )


class BulkChangeDifficultyIn(BaseModel):
    question_ids: List[int]
    difficulty: int = Field(..., ge=1, le=5)


class BulkAddTagsIn(BaseModel):
    question_ids: List[int]
    tags: List[str]


@router.post("/bulk-difficulty")
async def bulk_change_questions_difficulty(
    payload: BulkChangeDifficultyIn,
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
) -> dict:
    own_rows = (await session.execute(
        text("SELECT discipline_id FROM teacher_disciplines WHERE teacher_id = :tid").bindparams(tid=user.id)
    )).all()
    own_set = {int(r[0]) for r in own_rows}

    rows = (await session.execute(
        select(Question).where(
            Question.question_id.in_(payload.question_ids),
            Question.archived_at.is_(None)
        )
    )).scalars().all()

    if not rows:
        return {"ok": True, "count": 0}

    updated = 0
    for q in rows:
        if q.discipline_id not in own_set:
            continue
        q.difficulty = payload.difficulty
        await ensure_question_version(session, q.question_id, created_by_teacher_id=user.id)
        updated += 1

    if updated:
        await teacher_audit_log(
            session,
            teacher_id=user.id,
            action="bulk_difficulty_change",
            target_type="question",
            target_id=None,
            before={"count": updated},
            after={"question_ids": [q.question_id for q in rows if q.discipline_id in own_set], "difficulty": payload.difficulty},
        )
        await session.commit()

    return {"ok": True, "count": updated}


@router.post("/bulk-tags")
async def bulk_add_questions_tags(
    payload: BulkAddTagsIn,
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
) -> dict:
    own_rows = (await session.execute(
        text("SELECT discipline_id FROM teacher_disciplines WHERE teacher_id = :tid").bindparams(tid=user.id)
    )).all()
    own_set = {int(r[0]) for r in own_rows}

    rows = (await session.execute(
        select(Question).where(
            Question.question_id.in_(payload.question_ids),
            Question.archived_at.is_(None)
        )
    )).scalars().all()

    if not rows:
        return {"ok": True, "count": 0}

    tag_names = [t.strip() for t in payload.tags if t.strip()]
    if not tag_names:
        return {"ok": True, "count": 0}

    tags_map = {}
    for tname in tag_names:
        t = (await session.execute(
            select(QuestionTag).where(QuestionTag.name == tname)
        )).scalar_one_or_none()
        if not t:
            t = QuestionTag(name=tname)
            session.add(t)
            await session.flush()
        tags_map[tname] = t.tag_id

    updated = 0
    for q in rows:
        if q.discipline_id not in own_set:
            continue

        for tname, tag_id in tags_map.items():
            exists = (await session.execute(
                select(1).select_from(QuestionTagMap).where(
                    QuestionTagMap.question_id == q.question_id,
                    QuestionTagMap.tag_id == tag_id
                )
            )).scalar()
            if not exists:
                session.add(QuestionTagMap(question_id=q.question_id, tag_id=tag_id))

        await ensure_question_version(session, q.question_id, created_by_teacher_id=user.id)
        updated += 1

    if updated:
        await teacher_audit_log(
            session,
            teacher_id=user.id,
            action="bulk_tags_added",
            target_type="question",
            target_id=None,
            before={"count": updated},
            after={"question_ids": [q.question_id for q in rows if q.discipline_id in own_set], "tags": tag_names},
        )
        await session.commit()

    return {"ok": True, "count": updated}
