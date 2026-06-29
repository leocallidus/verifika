from __future__ import annotations

import io
from typing import List
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import CurrentUser, require_teacher
from app.db.models import (
    AnswerComment,
    AnswerOption,
    Discipline,
    Group,
    Question,
    Student,
    StudentAnswer,
    TeacherDiscipline,
    TestSession,
    TestSessionGradeOverride,
    DisciplineTopic,
)
from app.db.session import get_session
from app.services.notifications_helpers import record_event

try:
    from weasyprint import HTML as _WeasyHTML  # type: ignore
    _HAS_WEASY = True
except Exception:
    _WeasyHTML = None
    _HAS_WEASY = False

try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    _HAS_OPENPYXL = True
except Exception:
    _HAS_OPENPYXL = False

router = APIRouter(prefix="/teacher/reports", tags=["v2.reports"])


@router.get("/students/{student_id}.pdf")
async def student_pdf(
    student_id: int,
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
):
    student = (await session.execute(select(Student).where(Student.student_id == student_id))).scalar_one_or_none()
    if student is None:
        raise HTTPException(404, "student not found")
    rows = (await session.execute(
        select(TestSession, Discipline.name)
        .join(Discipline, Discipline.discipline_id == TestSession.discipline_id)
        .where(TestSession.student_id == student_id, TestSession.teacher_id == user.id)
        .order_by(TestSession.started_at.desc())
    )).all()
    grouped: dict[str, list[dict]] = {}
    for sess, dname in rows:
        grouped.setdefault(dname, [])
        answers = (await session.execute(
            select(Question.text, AnswerOption.text, AnswerOption.is_correct, AnswerOption.option_number, Question.question_id)
            .select_from(StudentAnswer)
            .join(Question, Question.question_id == StudentAnswer.question_id)
            .join(AnswerOption, AnswerOption.option_id == StudentAnswer.selected_option_id)
            .where(StudentAnswer.session_id == sess.session_id)
        )).all()
        comments = (await session.execute(
            select(AnswerComment.question_id, AnswerComment.body).where(
                AnswerComment.session_id == sess.session_id,
            )
        )).all()
        cmt_map = {qid: body for qid, body in comments}
        grouped[dname].append({
            "session_id": sess.session_id,
            "started_at": sess.started_at.isoformat(),
            "completed_at": sess.completed_at.isoformat() if sess.completed_at else "—",
            "score": sess.score, "max_score": sess.max_score,
            "answers": [
                {
                    "question_text": qt, "chosen_option_text": ot,
                    "is_correct": "✓" if ok else "✗",
                    "option_number": onum, "comment": cmt_map.get(qid),
                }
                for qt, ot, ok, onum, qid in answers
            ],
        })
    html = _render_student_html(
        full_name=f"{student.last_name} {student.first_name} {student.middle_name or ''}".strip(),
        email=student.email, grouped=grouped, user_name=f"{user.full_name}",
    )
    await record_event(session, "report_generated", target=f"student:{student_id}", metadata={"format": "pdf"})

    if _HAS_WEASY:
        pdf_bytes = _WeasyHTML(string=html, base_url=".").write_pdf()
        return StreamingResponse(
            iter([pdf_bytes]),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="student-{student_id}.pdf"'},
        )
    return StreamingResponse(
        iter([html.encode("utf-8")]),
        media_type="text/html; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="student-{student_id}.html"'},
    )


@router.get("/sessions/{session_id}.pdf")
async def session_pdf(
    session_id: int,
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
):
    sess = (await session.execute(
        select(TestSession)
        .where(TestSession.session_id == session_id, TestSession.teacher_id == user.id)
    )).scalar_one_or_none()
    if sess is None:
        raise HTTPException(404, "session not found")
        
    student = (await session.execute(
        select(Student).where(Student.student_id == sess.student_id)
    )).scalar_one_or_none()
    if student is None:
        raise HTTPException(404, "student not found")
        
    discipline = (await session.execute(
        select(Discipline).where(Discipline.discipline_id == sess.discipline_id)
    )).scalar_one_or_none()
    dname = discipline.name if discipline else "Дисциплина"
    
    answers = (await session.execute(
        select(Question.text, AnswerOption.text, AnswerOption.is_correct, AnswerOption.option_number, Question.question_id)
        .select_from(StudentAnswer)
        .join(Question, Question.question_id == StudentAnswer.question_id)
        .join(AnswerOption, AnswerOption.option_id == StudentAnswer.selected_option_id)
        .where(StudentAnswer.session_id == session_id)
    )).all()
    
    comments = (await session.execute(
        select(AnswerComment.question_id, AnswerComment.body).where(
            AnswerComment.session_id == session_id,
        )
    )).all()
    cmt_map = {qid: body for qid, body in comments}
    
    grouped = {
        dname: [{
            "session_id": sess.session_id,
            "started_at": sess.started_at.isoformat(),
            "completed_at": sess.completed_at.isoformat() if sess.completed_at else "—",
            "score": sess.score,
            "max_score": sess.max_score,
            "answers": [
                {
                    "question_text": qt, "chosen_option_text": ot,
                    "is_correct": "✓" if ok else "✗",
                    "option_number": onum, "comment": cmt_map.get(qid),
                }
                for qt, ot, ok, onum, qid in answers
            ],
        }]
    }
    
    html = _render_student_html(
        full_name=f"{student.last_name} {student.first_name} {student.middle_name or ''}".strip(),
        email=student.email, grouped=grouped, user_name=f"{user.full_name}",
    )
    
    await record_event(session, "report_generated", target=f"session:{session_id}", metadata={"format": "pdf"})
    
    if _HAS_WEASY:
        pdf_bytes = _WeasyHTML(string=html, base_url=".").write_pdf()
        return StreamingResponse(
            iter([pdf_bytes]),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="session-{session_id}.pdf"'},
        )
    return StreamingResponse(
        iter([html.encode("utf-8")]),
        media_type="text/html; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="session-{session_id}.html"'},
    )


@router.get("/groups/{group_id}/gradebook.xlsx")
async def group_gradebook(
    group_id: int,
    format: str = "xlsx",
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
):
    group = (await session.execute(select(Group).where(Group.group_id == group_id))).scalar_one_or_none()
    if group is None:
        raise HTTPException(404, "group not found")
    students = (await session.execute(
        select(Student).where(Student.group_id == group_id).order_by(Student.last_name)
    )).scalars().all()
    my_discs = (await session.execute(
        select(Discipline.discipline_id, Discipline.name)
        .join(TeacherDiscipline, TeacherDiscipline.discipline_id == Discipline.discipline_id)
        .where(TeacherDiscipline.teacher_id == user.id)
        .order_by(Discipline.name)
    )).all()
    disc_ids = [d for d, _ in my_discs]
    rows = (await session.execute(
        select(
            Student.student_id, TestSession.discipline_id, TestSession.session_id,
            TestSession.score, TestSession.max_score, TestSession.completed_at,
            TestSession.started_at, TestSessionGradeOverride.score,
            TestSession.status, TestSession.topic_id, TestSession.attempt_no,
        )
        .outerjoin(TestSession,
                   (TestSession.student_id == Student.student_id)
                   & (TestSession.discipline_id.in_(disc_ids))
                   & (TestSession.teacher_id == user.id))
        .outerjoin(TestSessionGradeOverride, TestSessionGradeOverride.session_id == TestSession.session_id)
        .where(Student.group_id == group_id)
    )).all()
    per_student: dict[int, dict[int, list]] = {s.student_id: {d: [] for d in disc_ids} for s in students}
    for sid, did, sess_id, score, maxs, compl, started, ovr_score, st_status, topic_id, att_no in rows:
        if sid in per_student and did in per_student[sid]:
            per_student[sid][did].append((sess_id, started, compl, score, maxs, ovr_score, st_status, topic_id, att_no))
    masked_email = lambda em: em[0] + "***" + em[em.index("@"):] if "@" in em else em

    await record_event(session, "report_generated", target=f"group:{group_id}/gradebook", metadata={"format": format})

    if format == "json":
        student_list = []
        for s in students:
            student_list.append({
                "student_id": s.student_id,
                "full_name": f"{s.last_name} {s.first_name} {s.middle_name or ''}".strip(),
                "email": masked_email(s.email),
            })
        
        discipline_list = [{"discipline_id": d_id, "name": d_name} for d_id, d_name in my_discs]
        
        topics = (await session.execute(
            select(DisciplineTopic.topic_id, DisciplineTopic.discipline_id, DisciplineTopic.name)
            .where(DisciplineTopic.discipline_id.in_(disc_ids))
            .where(DisciplineTopic.archived_at.is_(None))
            .order_by(DisciplineTopic.name)
        )).all()
        topic_list = [{"topic_id": t_id, "discipline_id": d_id, "name": t_name} for t_id, d_id, t_name in topics]
        
        session_list = []
        for sid, did, sess_id, score, maxs, compl, started, ovr_score, st_status, topic_id, att_no in rows:
            if sess_id is not None:
                session_list.append({
                    "student_id": sid,
                    "discipline_id": did,
                    "topic_id": topic_id,
                    "session_id": sess_id,
                    "score": score,
                    "max_score": maxs,
                    "started_at": started.isoformat() if started else None,
                    "completed_at": compl.isoformat() if compl else None,
                    "overridden_score": ovr_score,
                    "status": st_status,
                    "attempt_no": att_no,
                })
        return {
            "group_name": group.name,
            "students": student_list,
            "disciplines": discipline_list,
            "topics": topic_list,
            "sessions": session_list,
        }

    if format == "csv" or not _HAS_OPENPYXL:
        import csv
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow([f"Журнал: {group.name}", ""])
        header = ["ФИО", "email"] + [f"{n} (попытка #{i+1})" for _, n in my_discs for i in range(3)] + ["Лучший %", "Средний %"]
        w.writerow(header)
        for s in students:
            row = [f"{s.last_name} {s.first_name} {s.middle_name or ''}".strip(), masked_email(s.email)]
            for d_id, _ in my_discs:
                atts = per_student[s.student_id][d_id]
                for i in range(3):
                    if i < len(atts):
                        sess_id, started, compl, score, maxs, ovr, _st = atts[i]
                        eff = ovr if ovr is not None else (score or 0)
                        row.append(round(eff * 100.0 / maxs, 1) if maxs and compl else "—")
                    else:
                        row.append("")
            row += ["", ""]
            w.writerow(row)
        return StreamingResponse(
            iter([buf.getvalue().encode("utf-8")]),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="gradebook-{group.name}.csv"'},
        )

    wb = Workbook(); ws = wb.active; ws.title = "Журнал"
    ws.append([f"Журнал группы: {group.name}"])
    ws.append([f"Преподаватель: {user.full_name}"])
    ws.append([])
    disc_names = [n for _, n in my_discs]
    header = ["ФИО", "email"] + [f"{n} (п{i+1})" for n in disc_names for i in range(3)] + ["Лучший %", "Средний %"]
    ws.append(header)
    red_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    green_fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
    score_row_start = 5
    for s in students:
        row = [f"{s.last_name} {s.first_name} {s.middle_name or ''}".strip(), masked_email(s.email)]
        for d_id, _ in my_discs:
            for i in range(3):
                atts = per_student[s.student_id][d_id]
                if i < len(atts):
                    sess_id, started, compl, score, maxs, ovr, _st = atts[i]
                    eff = ovr if ovr is not None else (score or 0)
                    if maxs and compl:
                        row.append(round(eff * 100.0 / maxs, 1))
                    else:
                        row.append("")
                else:
                    row.append("")
        row += ["", ""]
        ws.append(row)
        r = ws.max_row
        start_col = 3
        end_col = start_col + 3 * len(disc_names) - 1
        if end_col >= start_col:
            cell = ws.cell(row=r, column=end_col + 1,
                           value=f'=IF(MAX({get_column_letter(start_col)}{r}:{get_column_letter(end_col)}{r})>0, MAX({get_column_letter(start_col)}{r}:{get_column_letter(end_col)}{r}), "—")')
            cell.alignment = cell.alignment.copy(horizontal="center")
            avg_cell = ws.cell(row=r, column=end_col + 2,
                               value=f'=IFERROR(ROUND(AVERAGEIF({get_column_letter(start_col)}{r}:{get_column_letter(end_col)}{r}, ">0"), 1), "—")')
            avg_cell.alignment = avg_cell.alignment.copy(horizontal="center")
            if isinstance(cell.value, (int, float)):
                avg_val = cell.value
                if isinstance(avg_val, (int, float)):
                    if avg_val < 50:
                        cell.fill = red_fill
                    elif avg_val < 70:
                        cell.fill = yellow_fill
                    else:
                        cell.fill = green_fill
        for c in range(start_col, end_col + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, (int, float)):
                if v < 50:
                    ws.cell(row=r, column=c).fill = red_fill
                elif v < 70:
                    ws.cell(row=r, column=c).fill = yellow_fill
                else:
                    ws.cell(row=r, column=c).fill = green_fill

    avg_col = ws.max_column
    data_start = score_row_start
    data_end = ws.max_row
    disc_avg_row = data_end + 3
    ws.cell(row=disc_avg_row, column=1, value="Средний % группы").font = Font(bold=True)
    for i, (d_id, d_name) in enumerate(my_discs):
        cells = [get_column_letter(3 + 3 * i + k) for k in range(3)]
        rng = ",".join(f'{c}{data_start}:{c}{data_end}' for c in cells)
        cell = ws.cell(
            row=disc_avg_row, column=3 + 3 * i,
            value=f'=IFERROR(ROUND(AVERAGE({rng}), 1), "—")',
        )
        cell.font = Font(bold=True)

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Средний % по дисциплине"
    chart.y_axis.title = "%"
    chart.x_axis.title = "Дисциплина"
    data_ref = Reference(ws, min_col=3, max_col=2 + 3 * len(disc_names) if disc_names else 2,
                         min_row=disc_avg_row, max_row=disc_avg_row)
    cats_ref = Reference(ws, min_col=3, max_col=2 + 3 * len(disc_names) if disc_names else 2,
                         min_row=4, max_row=4)
    chart.add_data(data_ref, titles_from_data=False)
    chart.set_categories(cats_ref)
    chart.width = 16
    chart.height = 8
    ws.add_chart(chart, f"A{disc_avg_row + 2}")

    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 12
        for r in range(1, ws.max_row + 1):
            v = ws.cell(row=r, column=col_idx).value
            if v:
                max_len = max(max_len, len(str(v)) + 2)
        ws.column_dimensions[letter].width = min(max_len, 40)

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    raw_fname = f"gradebook-g{group.group_id}.xlsx"
    encoded_fname_quoted = quote(f"gradebook-{group.name}.xlsx")
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{raw_fname}"; '
                f"filename*=UTF-8''{encoded_fname_quoted}"
            )
        },
    )


def _render_student_html(full_name: str, email: str, grouped: dict, user_name: str) -> str:
    parts = ["""<!DOCTYPE html><html lang="ru"><head><meta charset="utf-8"><style>
@page { size: A4; margin: 1.5cm }
body { font-family: 'DejaVu Sans', 'Helvetica', sans-serif; color: #1f2937; font-size: 11pt }
h1 { color: #4338ca; margin: 0 0 4pt 0; font-size: 18pt }
h2 { color: #4338ca; border-bottom: 1pt solid #c7d2fe; padding-bottom: 4pt; margin-top: 14pt }
.meta { color: #6b7280; font-size: 9pt; margin-bottom: 14pt }
table { width: 100%; border-collapse: collapse; margin-top: 4pt }
th, td { border-bottom: 0.5pt solid #e5e7eb; padding: 4pt 6pt; text-align: left; vertical-align: top }
th { background: #eef2ff; font-size: 9pt; color: #374151 }
.badge { display: inline-block; padding: 1pt 6pt; border-radius: 8pt; font-size: 9pt; font-weight: 600 }
.ok { background: #d1fae5; color: #065f46 }
.bad { background: #fee2e2; color: #991b1b }
.overridden { background: #fef3c7; color: #92400e }
.lvl { font-size: 9pt; color: #4b5563 }
.comment { background: #fef9c3; padding: 3pt 6pt; border-left: 3pt solid #facc15; font-size: 9pt; margin-top: 3pt }
""", "</style></head><body>"]
    parts.append(f"<h1>{full_name}</h1>")
    parts.append(f'<div class="meta">{email} · Преподаватель: {user_name} · Сформирован: автоматически</div>')
    for dname, sessions in grouped.items():
        parts.append(f'<h2>{dname}</h2>')
        if not sessions:
            parts.append('<p class="meta">Попыток нет.</p>')
            continue
        for s in sessions:
            parts.append(f'<div class="lvl">Сессия #{s["session_id"]} · {s["started_at"]} → {s["completed_at"]} · <b>{s["score"]}/{s["max_score"]}</b></div>')
            parts.append('<table><thead><tr><th style="width:55%">Вопрос</th><th>Ответ студента</th><th>Верно?</th></tr></thead><tbody>')
            for a in s["answers"]:
                badge_cls = "ok" if a["is_correct"] == "✓" else "bad"
                parts.append(f'<tr><td>{a["question_text"]}<br><span class="lvl">вариант №{a["option_number"]}</span></td><td>{a["chosen_option_text"]}')
                if a.get("comment"):
                    parts.append(f'<div class="comment">💬 {a["comment"]}</div>')
                parts.append(f'</td><td><span class="badge {badge_cls}">{a["is_correct"]}</span></td></tr>')
            parts.append('</tbody></table><br>')
    parts.append('</body></html>')
    return "".join(parts)


@router.get("/group/{group_id}.pdf")
async def group_pdf(
    group_id: int,
    scale: str = "5_point",
    user: CurrentUser = Depends(require_teacher),
    session: AsyncSession = Depends(get_session),
):
    group = (await session.execute(select(Group).where(Group.group_id == group_id))).scalar_one_or_none()
    if group is None:
        raise HTTPException(404, "group not found")
    students = (await session.execute(
        select(Student).where(Student.group_id == group_id).order_by(Student.last_name)
    )).scalars().all()
    my_discs = (await session.execute(
        select(Discipline.discipline_id, Discipline.name)
        .join(TeacherDiscipline, TeacherDiscipline.discipline_id == Discipline.discipline_id)
        .where(TeacherDiscipline.teacher_id == user.id)
        .order_by(Discipline.name)
    )).all()
    disc_ids = [d for d, _ in my_discs]
    rows = (await session.execute(
        select(
            Student.student_id, TestSession.discipline_id, TestSession.session_id,
            TestSession.score, TestSession.max_score, TestSession.completed_at,
            TestSession.started_at, TestSessionGradeOverride.score,
        )
        .outerjoin(TestSession,
                   (TestSession.student_id == Student.student_id)
                   & (TestSession.discipline_id.in_(disc_ids))
                   & (TestSession.teacher_id == user.id))
        .outerjoin(TestSessionGradeOverride, TestSessionGradeOverride.session_id == TestSession.session_id)
        .where(Student.group_id == group_id)
    )).all()
    per_student = {s.student_id: {d: [] for d in disc_ids} for s in students}
    for sid, did, sess_id, score, maxs, compl, started, ovr_score in rows:
        if sid in per_student and did in per_student[sid]:
            per_student[sid][did].append((sess_id, started, compl, score, maxs, ovr_score))
            
    html = _render_group_html(
        group_name=group.name,
        teacher_name=user.full_name,
        students=students,
        my_discs=my_discs,
        per_student=per_student,
        scale=scale,
    )
    
    await record_event(session, "report_generated", target=f"group:{group_id}/gradebook", metadata={"format": "pdf"})
    
    if _HAS_WEASY:
        pdf_bytes = _WeasyHTML(string=html, base_url=".").write_pdf()
        return StreamingResponse(
            iter([pdf_bytes]),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="group-{group_id}.pdf"'},
        )
    return StreamingResponse(
        iter([html.encode("utf-8")]),
        media_type="text/html; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="group-{group_id}.html"'},
    )


def _render_group_html(group_name: str, teacher_name: str, students: list, my_discs: list, per_student: dict, scale: str = "5_point") -> str:
    page_width = max(297, 110 + len(my_discs) * 30)
    font_size = "9pt" if len(my_discs) <= 4 else "8pt" if len(my_discs) <= 8 else "7pt"

    parts = [f"""<!DOCTYPE html><html lang="ru"><head><meta charset="utf-8"><style>
@page {{ size: {page_width}mm 210mm; margin: 1cm }}
body {{ font-family: 'DejaVu Sans', 'Helvetica', sans-serif; color: #1f2937; font-size: {font_size} }}
h1 {{ color: #4338ca; margin: 0 0 4pt 0; font-size: 16pt }}
.meta {{ color: #6b7280; font-size: 8pt; margin-bottom: 10pt }}
table {{ width: 100%; border-collapse: collapse; margin-top: 4pt; font-size: 8pt }}
th, td {{ border: 0.5pt solid #d1d5db; padding: 4pt; text-align: left; vertical-align: middle }}
th {{ background: #eef2ff; color: #374151; font-weight: bold }}
.avg-green {{ background: #d1fae5; color: #065f46; font-weight: bold }}
.avg-yellow {{ background: #fef3c7; color: #92400e; font-weight: bold }}
.avg-red {{ background: #fee2e2; color: #991b1b; font-weight: bold }}
</style></head><body>"""]
    parts.append(f"<h1>Журнал группы: {group_name}</h1>")
    parts.append(f'<div class="meta">Преподаватель: {teacher_name} · Сформирован: автоматически</div>')
    
    parts.append('<table><thead><tr><th>Студент</th><th>Email</th>')
    for _, dname in my_discs:
        parts.append(f'<th colspan="3">{dname}</th>')
        
    best_header = "Лучший %" if scale == "percent" else "Лучший балл"
    avg_header = "Средний %" if scale == "percent" else "Средний балл"
    parts.append(f'<th>{best_header}</th><th>{avg_header}</th></tr>')
    
    parts.append('<tr><th></th><th></th>')
    for _ in my_discs:
        parts.append('<th>Поп. 1</th><th>Поп. 2</th><th>Поп. 3</th>')
    parts.append('<th></th><th></th></tr></thead><tbody>')
    
    masked_email = lambda em: em[0] + "***" + em[em.index("@"):] if "@" in em else em
    
    from app.services.grade_calculator import FIVE_POINT_THRESHOLDS, TEN_POINT_THRESHOLDS

    def format_grade_short_py(percent: float, scale_val: str) -> str:
        pct = round(percent)
        if scale_val == "5_point":
            val = 1
            for min_pct, grade in FIVE_POINT_THRESHOLDS:
                if pct >= min_pct:
                    val = grade
                    break
            return str(val)
        if scale_val == "10_point":
            val = 1
            for min_pct, grade in TEN_POINT_THRESHOLDS:
                if pct >= min_pct:
                    val = grade
                    break
            return str(val)
        return f"{round(percent, 1)}%"

    for s in students:
        s_name = f"{s.last_name} {s.first_name} {s.middle_name or ''}".strip()
        parts.append(f'<tr><td><b>{s_name}</b></td><td>{masked_email(s.email)}</td>')
        
        all_percentages = []
        best_perc = 0.0
        
        for d_id, _ in my_discs:
            atts = per_student[s.student_id][d_id]
            for i in range(3):
                if i < len(atts):
                    sess_id, started, compl, score, maxs, ovr = atts[i]
                    eff = ovr if ovr is not None else (score or 0)
                    perc = round(eff * 100.0 / maxs, 1) if maxs and compl else 0.0
                    all_percentages.append(perc)
                    if perc > best_perc:
                        best_perc = perc
                    parts.append(f'<td>{format_grade_short_py(perc, scale)}</td>' if compl else '<td>—</td>')
                else:
                    parts.append('<td></td>')
                    
        avg_perc = round(sum(all_percentages) / len(all_percentages), 1) if all_percentages else 0.0
        
        best_cls = "avg-red" if best_perc < 50 else "avg-yellow" if best_perc < 75 else "avg-green"
        avg_cls = "avg-red" if avg_perc < 50 else "avg-yellow" if avg_perc < 75 else "avg-green"
        
        parts.append(f'<td class="{best_cls}">{format_grade_short_py(best_perc, scale)}</td>')
        parts.append(f'<td class="{avg_cls}">{format_grade_short_py(avg_perc, scale)}</td></tr>')
        
    parts.append('</tbody></table></body></html>')
    return "".join(parts)
