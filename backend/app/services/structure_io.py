from __future__ import annotations

import csv
import io
import secrets
import string
from dataclasses import dataclass, field
from typing import Any, Optional

from fastapi import HTTPException, UploadFile
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import hash_password
from app.db.models import (
    AnswerOption,
    Discipline,
    DisciplineTopic,
    Group,
    Question,
    Student,
    Teacher,
    TeacherDiscipline,
    UserCredential,
)


try:
    from openpyxl import Workbook, load_workbook
except Exception:  # pragma: no cover - dependency exists in uv env, fallback is for runtime safety.
    Workbook = None
    load_workbook = None


STRUCTURE_COLUMNS = [
    "row_type",
    "external_id",
    "name",
    "description",
    "admission_year",
    "last_name",
    "first_name",
    "middle_name",
    "email",
    "login",
    "group_name",
    "initial_password",
    "discipline_name",
    "teacher_login",
    "topic_name",
    "sort_order",
    "question_external_id",
    "question_text",
    "qtype",
    "difficulty",
    "short_pattern",
    "numeric_tolerance",
    "correct_bool",
    "explanation",
    "option_number",
    "option_text",
    "is_correct",
    "match_left",
    "match_right",
    "correct_position",
]

XLSX_SHEETS: dict[str, list[str]] = {
    "groups": ["name", "admission_year"],
    "students": ["last_name", "first_name", "middle_name", "email", "login", "group_name", "initial_password"],
    "disciplines": ["name", "description", "teacher_login"],
    "topics": ["discipline_name", "name", "description", "sort_order"],
    "questions": [
        "external_id", "discipline_name", "topic_name", "question_text", "qtype", "difficulty",
        "short_pattern", "numeric_tolerance", "correct_bool", "explanation",
    ],
    "options": [
        "question_external_id", "question_text", "option_number", "option_text", "is_correct",
        "match_left", "match_right", "correct_position",
    ],
}

ROW_TYPE_BY_SHEET = {
    "groups": "group",
    "students": "student",
    "disciplines": "discipline",
    "topics": "topic",
    "questions": "question",
    "options": "option",
}

ALLOWED_QTYPES = {"single", "multi", "short", "numeric", "match", "text", "order", "bool", "cloze"}


@dataclass
class StructureImportResult:
    created: dict[str, int] = field(default_factory=lambda: {
        "groups": 0,
        "students": 0,
        "disciplines": 0,
        "topics": 0,
        "questions": 0,
        "options": 0,
    })
    updated: dict[str, int] = field(default_factory=lambda: {
        "groups": 0,
        "students": 0,
        "disciplines": 0,
        "topics": 0,
        "questions": 0,
        "options": 0,
    })
    skipped: int = 0
    errors: list[dict[str, Any]] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {
            "created": self.created,
            "updated": self.updated,
            "skipped": self.skipped,
            "errors": self.errors,
        }


def _s(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _int(value: Any, default: Optional[int] = None) -> Optional[int]:
    raw = _s(value)
    if raw == "":
        return default
    return int(float(raw))


def _bool(value: Any, default: bool = False) -> bool:
    raw = _s(value).lower()
    if raw == "":
        return default
    return raw in {"1", "true", "yes", "y", "да", "истина"}


def _float(value: Any) -> Optional[float]:
    raw = _s(value)
    if raw == "":
        return None
    return float(raw.replace(",", "."))


def _login_from_email(email: str) -> str:
    base = (email.split("@", 1)[0] or "student").lower()
    allowed = string.ascii_lowercase + string.digits + "._-"
    cleaned = "".join(ch for ch in base if ch in allowed).strip("._-")
    return (cleaned or "student")[:48]


async def _unique_student_login(session: AsyncSession, preferred: str) -> str:
    candidate = preferred[:64]
    exists = (await session.execute(
        select(Student.student_id).where(Student.login == candidate)
    )).scalar_one_or_none()
    if exists is None:
        return candidate
    for _ in range(20):
        suffix = secrets.token_hex(3)
        candidate = f"{preferred[:56]}_{suffix}"
        exists = (await session.execute(
            select(Student.student_id).where(Student.login == candidate)
        )).scalar_one_or_none()
        if exists is None:
            return candidate
    raise ValueError("cannot generate unique student login")


async def _read_rows(file: UploadFile) -> list[dict[str, Any]]:
    raw = await file.read()
    name = (file.filename or "").lower()
    if name.endswith(".xlsx"):
        if load_workbook is None:
            raise HTTPException(status_code=409, detail="xlsx support is not available")
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        rows: list[dict[str, Any]] = []
        for sheet_name, row_type in ROW_TYPE_BY_SHEET.items():
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            iterator = sheet.iter_rows(values_only=True)
            headers = [_s(cell).lower() for cell in next(iterator, [])]
            if not headers:
                continue
            for index, values in enumerate(iterator, start=2):
                row = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))}
                if not any(_s(v) for v in row.values()):
                    continue
                row["row_type"] = row_type
                row["_row_number"] = index
                row["_sheet"] = sheet_name
                rows.append(row)
        return rows

    text = raw.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    rows = []
    for index, row in enumerate(reader, start=2):
        if not any(_s(v) for v in row.values()):
            continue
        clean = {(_s(k).lower()): v for k, v in row.items() if k is not None}
        clean["_row_number"] = index
        clean["_sheet"] = "csv"
        rows.append(clean)
    return rows


async def _get_or_create_group(
    session: AsyncSession,
    name: str,
    admission_year: Optional[int],
    result: StructureImportResult,
) -> Group:
    group = (await session.execute(select(Group).where(Group.name == name))).scalar_one_or_none()
    if group is not None:
        if admission_year is not None and group.admission_year != admission_year:
            group.admission_year = admission_year
            result.updated["groups"] += 1
        else:
            result.skipped += 1
        return group
    group = Group(name=name, admission_year=admission_year or 2026)
    session.add(group)
    await session.flush()
    result.created["groups"] += 1
    return group


async def _get_or_create_discipline(
    session: AsyncSession,
    name: str,
    description: str,
    teacher_id: Optional[int],
    result: StructureImportResult,
) -> Discipline:
    discipline = (await session.execute(select(Discipline).where(Discipline.name == name))).scalar_one_or_none()
    if discipline is not None:
        changed = False
        if description and discipline.description != description:
            discipline.description = description
            changed = True
        if changed:
            result.updated["disciplines"] += 1
        else:
            result.skipped += 1
    else:
        discipline = Discipline(name=name, description=description or None, created_by=teacher_id)
        session.add(discipline)
        await session.flush()
        result.created["disciplines"] += 1
    if teacher_id is not None:
        link = (await session.execute(
            select(TeacherDiscipline).where(
                TeacherDiscipline.teacher_id == teacher_id,
                TeacherDiscipline.discipline_id == discipline.discipline_id,
            )
        )).scalar_one_or_none()
        if link is None:
            session.add(TeacherDiscipline(teacher_id=teacher_id, discipline_id=discipline.discipline_id))
    return discipline


async def _discipline_by_name(session: AsyncSession, name: str, teacher_id: Optional[int]) -> Discipline:
    stmt = select(Discipline).where(Discipline.name == name)
    if teacher_id is not None:
        stmt = stmt.join(TeacherDiscipline, TeacherDiscipline.discipline_id == Discipline.discipline_id).where(
            TeacherDiscipline.teacher_id == teacher_id
        )
    discipline = (await session.execute(stmt)).scalar_one_or_none()
    if discipline is None:
        raise ValueError(f"discipline not found: {name}")
    return discipline


async def _topic_by_name(
    session: AsyncSession,
    discipline_id: int,
    name: str,
) -> Optional[DisciplineTopic]:
    return (await session.execute(
        select(DisciplineTopic).where(
            DisciplineTopic.discipline_id == discipline_id,
            DisciplineTopic.name == name,
        )
    )).scalar_one_or_none()


async def import_structure(
    session: AsyncSession,
    file: UploadFile,
    *,
    actor_teacher_id: Optional[int] = None,
) -> StructureImportResult:
    rows = await _read_rows(file)
    result = StructureImportResult()
    question_by_external: dict[str, Question] = {}
    question_by_text: dict[str, Question] = {}

    for wanted in ("group", "discipline", "student", "topic", "question", "option"):
        for row in [item for item in rows if _s(item.get("row_type")).lower() == wanted]:
            row_number = int(row.get("_row_number") or 0)
            try:
                if wanted == "group":
                    name = _s(row.get("name") or row.get("group_name"))
                    if not name:
                        raise ValueError("group name is required")
                    await _get_or_create_group(session, name, _int(row.get("admission_year")), result)

                elif wanted == "discipline":
                    name = _s(row.get("name") or row.get("discipline_name"))
                    if not name:
                        raise ValueError("discipline name is required")
                    teacher_id = actor_teacher_id
                    teacher_login = _s(row.get("teacher_login"))
                    if teacher_id is None and teacher_login:
                        teacher = (await session.execute(
                            select(Teacher).where(Teacher.login == teacher_login)
                        )).scalar_one_or_none()
                        if teacher is None:
                            raise ValueError(f"teacher not found: {teacher_login}")
                        teacher_id = teacher.teacher_id
                    await _get_or_create_discipline(session, name, _s(row.get("description")), teacher_id, result)

                elif wanted == "student":
                    group_name = _s(row.get("group_name"))
                    if not group_name:
                        raise ValueError("group_name is required")
                    group = await _get_or_create_group(session, group_name, None, result)
                    email = _s(row.get("email")).lower()
                    if not email:
                        raise ValueError("email is required")
                    student = (await session.execute(
                        select(Student).where(or_(Student.email == email, Student.login == _s(row.get("login")).lower()))
                    )).scalar_one_or_none()
                    if student is not None:
                        changed = False
                        if student.group_id != group.group_id:
                            student.group_id = group.group_id
                            changed = True
                        for attr in ("last_name", "first_name", "middle_name"):
                            value = _s(row.get(attr))
                            if value and getattr(student, attr) != value:
                                setattr(student, attr, value)
                                changed = True
                        result.updated["students"] += 1 if changed else 0
                        result.skipped += 0 if changed else 1
                        continue
                    preferred_login = _s(row.get("login")).lower() or _login_from_email(email)
                    login = await _unique_student_login(session, preferred_login)
                    student = Student(
                        last_name=_s(row.get("last_name")) or "Student",
                        first_name=_s(row.get("first_name")) or "Imported",
                        middle_name=_s(row.get("middle_name")) or None,
                        email=email,
                        login=login,
                        group_id=group.group_id,
                    )
                    session.add(student)
                    await session.flush()
                    password = _s(row.get("initial_password")) or secrets.token_urlsafe(10)
                    session.add(UserCredential(
                        role="student",
                        user_id=student.student_id,
                        password_hash=hash_password(password),
                    ))
                    result.created["students"] += 1

                elif wanted == "topic":
                    discipline = await _discipline_by_name(session, _s(row.get("discipline_name")), actor_teacher_id)
                    name = _s(row.get("name") or row.get("topic_name"))
                    if not name:
                        raise ValueError("topic name is required")
                    topic = await _topic_by_name(session, discipline.discipline_id, name)
                    if topic is not None:
                        changed = False
                        description = _s(row.get("description"))
                        sort_order = _int(row.get("sort_order"))
                        if description and topic.description != description:
                            topic.description = description
                            changed = True
                        if sort_order is not None and topic.sort_order != sort_order:
                            topic.sort_order = sort_order
                            changed = True
                        result.updated["topics"] += 1 if changed else 0
                        result.skipped += 0 if changed else 1
                        continue
                    session.add(DisciplineTopic(
                        discipline_id=discipline.discipline_id,
                        name=name,
                        description=_s(row.get("description")) or None,
                        sort_order=_int(row.get("sort_order"), 100) or 100,
                        created_by=actor_teacher_id,
                    ))
                    result.created["topics"] += 1

                elif wanted == "question":
                    discipline = await _discipline_by_name(session, _s(row.get("discipline_name")), actor_teacher_id)
                    topic_name = _s(row.get("topic_name"))
                    topic_id = None
                    if topic_name:
                        topic = await _topic_by_name(session, discipline.discipline_id, topic_name)
                        if topic is None:
                            raise ValueError(f"topic not found: {topic_name}")
                        topic_id = topic.topic_id
                    text_value = _s(row.get("question_text") or row.get("text"))
                    if not text_value:
                        raise ValueError("question_text is required")
                    existing = (await session.execute(
                        select(Question).where(
                            Question.discipline_id == discipline.discipline_id,
                            Question.topic_id == topic_id,
                            Question.text == text_value,
                        )
                    )).scalar_one_or_none()
                    if existing is not None:
                        result.skipped += 1
                        q = existing
                    else:
                        qtype = _s(row.get("qtype")) or "single"
                        if qtype not in ALLOWED_QTYPES:
                            raise ValueError(f"invalid qtype: {qtype}")
                        q = Question(
                            discipline_id=discipline.discipline_id,
                            topic_id=topic_id,
                            text=text_value,
                            qtype=qtype,
                            difficulty=_int(row.get("difficulty"), 1) or 1,
                            short_pattern=_s(row.get("short_pattern")) or None,
                            numeric_tolerance=_float(row.get("numeric_tolerance")),
                            correct_bool=_bool(row.get("correct_bool")) if _s(row.get("correct_bool")) else None,
                            explanation=_s(row.get("explanation")) or None,
                        )
                        session.add(q)
                        await session.flush()
                        result.created["questions"] += 1
                    external_id = _s(row.get("external_id") or row.get("question_external_id"))
                    if external_id:
                        question_by_external[external_id] = q
                    question_by_text[text_value] = q

                elif wanted == "option":
                    external_id = _s(row.get("question_external_id"))
                    question_text = _s(row.get("question_text"))
                    q = question_by_external.get(external_id) if external_id else None
                    q = q or question_by_text.get(question_text)
                    if q is None and question_text:
                        q = (await session.execute(select(Question).where(Question.text == question_text))).scalars().first()
                    if q is None:
                        raise ValueError("question not found for option")
                    option_number = _int(row.get("option_number"))
                    option_text = _s(row.get("option_text") or row.get("text"))
                    if option_number is None or option_number < 1:
                        raise ValueError("option_number is required")
                    if not option_text:
                        raise ValueError("option_text is required")
                    existing = (await session.execute(
                        select(AnswerOption).where(
                            AnswerOption.question_id == q.question_id,
                            AnswerOption.option_number == option_number,
                        )
                    )).scalar_one_or_none()
                    if existing is not None:
                        changed = False
                        if existing.text != option_text:
                            existing.text = option_text
                            changed = True
                        correct = _bool(row.get("is_correct"))
                        if existing.is_correct != correct:
                            existing.is_correct = correct
                            changed = True
                        result.updated["options"] += 1 if changed else 0
                        result.skipped += 0 if changed else 1
                        continue
                    session.add(AnswerOption(
                        question_id=q.question_id,
                        option_number=option_number,
                        text=option_text,
                        is_correct=_bool(row.get("is_correct")),
                        match_left=_s(row.get("match_left")) or None,
                        match_right=_s(row.get("match_right")) or None,
                        correct_position=_int(row.get("correct_position")),
                    ))
                    result.created["options"] += 1
            except Exception as exc:
                result.errors.append({
                    "row": row_number,
                    "type": wanted,
                    "message": str(exc),
                })

    await session.commit()
    return result


async def _teacher_discipline_ids(session: AsyncSession, teacher_id: Optional[int]) -> Optional[list[int]]:
    if teacher_id is None:
        return None
    return [
        int(value)
        for value in (await session.execute(
            select(TeacherDiscipline.discipline_id).where(TeacherDiscipline.teacher_id == teacher_id)
        )).scalars().all()
    ]


async def export_structure_rows(
    session: AsyncSession,
    *,
    teacher_id: Optional[int] = None,
    include_archived: bool = False,
) -> dict[str, list[dict[str, Any]]]:
    discipline_ids = await _teacher_discipline_ids(session, teacher_id)
    discipline_filter = []
    if discipline_ids is not None:
        discipline_filter.append(Discipline.discipline_id.in_(discipline_ids or [-1]))
    if not include_archived:
        discipline_filter.append(Discipline.archived_at.is_(None))

    groups = (await session.execute(
        select(Group).where(Group.archived_at.is_(None) if not include_archived else True).order_by(Group.name)
    )).scalars().all()
    students = (await session.execute(
        select(Student, Group)
        .join(Group, Group.group_id == Student.group_id)
        .where(Student.archived_at.is_(None) if not include_archived else True)
        .order_by(Group.name, Student.last_name, Student.first_name)
    )).all()
    disciplines = (await session.execute(
        select(Discipline).where(*discipline_filter).order_by(Discipline.name)
    )).scalars().all()
    selected_ids = [d.discipline_id for d in disciplines]
    topics = (await session.execute(
        select(DisciplineTopic, Discipline)
        .join(Discipline, Discipline.discipline_id == DisciplineTopic.discipline_id)
        .where(
            DisciplineTopic.discipline_id.in_(selected_ids or [-1]),
            DisciplineTopic.archived_at.is_(None) if not include_archived else True,
        )
        .order_by(Discipline.name, DisciplineTopic.sort_order, DisciplineTopic.name)
    )).all()
    questions = (await session.execute(
        select(Question, Discipline, DisciplineTopic)
        .join(Discipline, Discipline.discipline_id == Question.discipline_id)
        .outerjoin(DisciplineTopic, DisciplineTopic.topic_id == Question.topic_id)
        .where(
            Question.discipline_id.in_(selected_ids or [-1]),
            Question.archived_at.is_(None) if not include_archived else True,
        )
        .order_by(Discipline.name, DisciplineTopic.sort_order.nullslast(), Question.question_id)
    )).all()
    question_ids = [q.question_id for q, _discipline, _topic in questions]
    options = (await session.execute(
        select(AnswerOption, Question)
        .join(Question, Question.question_id == AnswerOption.question_id)
        .where(AnswerOption.question_id.in_(question_ids or [-1]))
        .order_by(AnswerOption.question_id, AnswerOption.option_number)
    )).all()

    return {
        "groups": [
            {"name": g.name, "admission_year": g.admission_year}
            for g in groups
        ],
        "students": [
            {
                "last_name": s.last_name,
                "first_name": s.first_name,
                "middle_name": s.middle_name or "",
                "email": s.email,
                "login": s.login,
                "group_name": g.name,
                "initial_password": "",
            }
            for s, g in students
        ],
        "disciplines": [
            {
                "name": d.name,
                "description": d.description or "",
                "teacher_login": "",
            }
            for d in disciplines
        ],
        "topics": [
            {
                "discipline_name": d.name,
                "name": t.name,
                "description": t.description or "",
                "sort_order": t.sort_order,
            }
            for t, d in topics
        ],
        "questions": [
            {
                "external_id": f"q{q.question_id}",
                "discipline_name": d.name,
                "topic_name": t.name if t is not None else "",
                "question_text": q.text,
                "qtype": q.qtype,
                "difficulty": q.difficulty,
                "short_pattern": q.short_pattern or "",
                "numeric_tolerance": q.numeric_tolerance if q.numeric_tolerance is not None else "",
                "correct_bool": "" if q.correct_bool is None else int(q.correct_bool),
                "explanation": q.explanation or "",
            }
            for q, d, t in questions
        ],
        "options": [
            {
                "question_external_id": f"q{q.question_id}",
                "question_text": q.text,
                "option_number": option.option_number,
                "option_text": option.text,
                "is_correct": int(bool(option.is_correct)),
                "match_left": option.match_left or "",
                "match_right": option.match_right or "",
                "correct_position": option.correct_position or "",
            }
            for option, q in options
        ],
    }


def structure_csv_response(rows_by_sheet: dict[str, list[dict[str, Any]]]) -> bytes:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=STRUCTURE_COLUMNS, extrasaction="ignore")
    writer.writeheader()
    for sheet, row_type in ROW_TYPE_BY_SHEET.items():
        for row in rows_by_sheet.get(sheet, []):
            full = {column: "" for column in STRUCTURE_COLUMNS}
            full.update(row)
            full["row_type"] = row_type
            writer.writerow(full)
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def structure_xlsx_response(rows_by_sheet: dict[str, list[dict[str, Any]]]) -> bytes:
    if Workbook is None:
        raise HTTPException(status_code=409, detail="xlsx support is not available")
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    for sheet_name, columns in XLSX_SHEETS.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(columns)
        for row in rows_by_sheet.get(sheet_name, []):
            sheet.append([row.get(column, "") for column in columns])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
